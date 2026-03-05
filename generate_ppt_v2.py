
import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    try:
        prs = Presentation()
    except Exception as e:
        print(f"Error initializing Presentation: {e}")
        return

    # --- Helper Functions ---
    def add_title_slide(title_text, subtitle_text):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = title_text
        subtitle.text = subtitle_text

    def add_bullet_slide(title_text, content_list):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        body = slide.placeholders[1]
        title.text = title_text
        tf = body.text_frame
        tf.word_wrap = True
        for item in content_list:
            p = tf.add_paragraph()
            p.text = item
            p.level = 0

    def add_quadrant_slide(title_text, quadrants):
        # Simulating the quadrant layout from the example
        slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title Only
        slide.shapes.title.text = title_text
        
        # approximate positions
        width = prs.slide_width
        height = prs.slide_height
        mid_x = width / 2
        mid_y = height / 2 + Inches(0.5)
        
        # Top Left
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), mid_x - Inches(1), Inches(2.5))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = quadrants[0]['title']
        p.font.bold = True
        p.font.size = Pt(20)
        p = tf.add_paragraph()
        p.text = quadrants[0]['content']
        p.font.size = Pt(14)

        # Top Right
        txBox = slide.shapes.add_textbox(mid_x + Inches(0.5), Inches(1.5), mid_x - Inches(1), Inches(2.5))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = quadrants[1]['title']
        p.font.bold = True
        p.font.size = Pt(20)
        p = tf.add_paragraph()
        p.text = quadrants[1]['content']
        p.font.size = Pt(14)

        # Bottom Left
        txBox = slide.shapes.add_textbox(Inches(0.5), mid_y, mid_x - Inches(1), Inches(2.5))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = quadrants[2]['title']
        p.font.bold = True
        p.font.size = Pt(20)
        p = tf.add_paragraph()
        p.text = quadrants[2]['content']
        p.font.size = Pt(14)

        # Bottom Right
        txBox = slide.shapes.add_textbox(mid_x + Inches(0.5), mid_y, mid_x - Inches(1), Inches(2.5))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = quadrants[3]['title']
        p.font.bold = True
        p.font.size = Pt(20)
        p = tf.add_paragraph()
        p.text = quadrants[3]['content']
        p.font.size = Pt(14)
        
        # Draw dividing lines
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, mid_x, Inches(1.5), Inches(0.05), height - Inches(2))
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), mid_y - Inches(0.2), width - Inches(1), Inches(0.05))


    def add_table_slide(title_text, tables_data):
        slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title Only
        slide.shapes.title.text = title_text
        
        # Determine layout based on number of tables (1 or 2 side by side)
        cnt = len(tables_data)
        
        for i, table_info in enumerate(tables_data):
            rows = len(table_info['data']) + 1
            cols = len(table_info['headers'])
            
            # Positioning
            if cnt == 1:
                left = Inches(1)
                top = Inches(2)
                width = Inches(8)
            else:
                width = Inches(4.5)
                top = Inches(2)
                if i == 0:
                    left = Inches(0.5)
                else:
                    left = Inches(5.2)

            # Table Title for multiple tables
            if cnt > 1:
                txBox = slide.shapes.add_textbox(left, top - Inches(0.4), width, Inches(0.4))
                txBox.text_frame.text = f"Table Name: {table_info['name']}"
                txBox.text_frame.paragraphs[0].font.bold = True

            table_shape = slide.shapes.add_table(rows, cols, left, top, width, Inches(0.3 * rows)).table
            
            # Headers
            for idx, header in enumerate(table_info['headers']):
                cell = table_shape.cell(0, idx)
                cell.text = header
                cell.text_frame.paragraphs[0].font.bold = True
                
            # Data
            for r_idx, row_data in enumerate(table_info['data']):
                for c_idx, cell_data in enumerate(row_data):
                    table_shape.cell(r_idx + 1, c_idx).text = str(cell_data)

    def add_placeholder_slide(title_text, placeholder_text):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title_text
        txBox = slide.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(2))
        p = txBox.text_frame.add_paragraph()
        p.text = placeholder_text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(24)
        
        # Add a placeholder rectangle
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(8), Inches(4.5)).text = placeholder_text

    # --- Content Generation ---

    # 1. Title Slide
    add_title_slide("Dermascan AI: AI Skin Expert", "Final Year Project Presentation\n\nPresented By\n[Your Name]\n\nDepartment\nComputer Science")

    # 2. Abstract
    add_bullet_slide("Abstract", [
        "Purpose: A web-based platform revolutionizing skin disease diagnosis by connecting patients and dermatologists.",
        "Direct Interaction: Eliminates barriers between patients and doctors, ensuring timely medical advice.",
        "AI-Powered Diagnosis: Uses Deep Learning (TensorFlow/Keras) for instant preliminary analysis of skin lesions.",
        "Streamlined Management: Centralized system for appointments, prescriptions, and medical history."
    ])

    # 3. Existing System
    existing_quadrants = [
        {"title": "Limited Accessibility", "content": "Patients in remote/rural areas struggle to access specialized dermatological care without traveling long distances."},
        {"title": "Manual Diagnosis", "content": "Diagnosis relies solely on physical visits, which can be time-consuming and delayed due to doctor unavailability."},
        {"title": "Inefficient Appointments", "content": "Traditional booking leads to long waiting times, scheduling conflicts, and lack of transparency."},
        {"title": "Physical Records", "content": "Medical history is often maintained in paper files, making it hard to track disease progression over time."}
    ]
    add_quadrant_slide("Existing System", existing_quadrants)

    # 4. Proposed System
    proposed_quadrants = [
        {"title": "Unified Digital Platform", "content": "A centralized web portal connecting Patients, Doctors, and Admins for seamless healthcare delivery."},
        {"title": "AI & Automation", "content": "Integrated AI model for instant skin disease detection and probability scoring (~95% accuracy)."},
        {"title": "Dynamic Scheduling", "content": "Real-time appointment booking system with availability management for doctors."},
        {"title": "Digital Scan & Appointment History", "content": "Secure cloud-based storage of AI analysis results, scan images, and appointment records for easy access."}
    ]
    add_quadrant_slide("Proposed System", proposed_quadrants)

    # 5. Tables
    # We will create grouped table slides
    
    # Slide 5: User & Profile
    tables_s5 = [
        {
            "name": "User (Django Auth)",
            "headers": ["Field", "Type", "Size", "Constraint"],
            "data": [
                ["id", "INT", "11", "PK, Auto"],
                ["username", "VARCHAR", "150", "Unique"],
                ["password", "VARCHAR", "128", "Not Null"],
                ["email", "VARCHAR", "254", "Nullable"]
            ]
        },
        {
            "name": "User_profile",
            "headers": ["Field", "Type", "Size", "Constraint"],
            "data": [
                ["USER_ID", "INT", "11", "FK (User)"],
                ["name", "VARCHAR", "100", "Not Null"],
                ["gender", "VARCHAR", "10", "Not Null"],
                ["phone", "VARCHAR", "15", "Not Null"]
            ]
        }
    ]
    add_table_slide("Tables: Users", tables_s5)

    # Slide 6: Doctor & Schedule
    tables_s6 = [
        {
            "name": "Doctor",
            "headers": ["Field", "Type", "Constraint"],
            "data": [
                ["USER_ID", "INT", "FK"],
                ["regno", "VARCHAR", "Unique"],
                ["smc", "VARCHAR", "Not Null"],
                ["status", "VARCHAR", "Pending"]
            ]
        },
        {
            "name": "Doctor_schedule",
            "headers": ["Field", "Type", "Constraint"],
            "data": [
                ["DOCTOR_ID", "INT", "FK"],
                ["date", "VARCHAR", "Not Null"],
                ["start/end_time", "VARCHAR", "Not Null"],
                ["fees", "DECIMAL", "Not Null"]
            ]
        }
    ]
    add_table_slide("Tables: Doctors", tables_s6)

    # Slide 7: Bookings & Payments
    tables_s7 = [
         {
            "name": "Booking",
            "headers": ["Field", "Type", "Constraint"],
            "data": [
                ["PATIENT_ID", "INT", "FK"],
                ["schedule_id", "INT", "FK"],
                ["date", "VARCHAR", "Not Null"],
                ["status", "VARCHAR", "Pending"]
            ]
        },
        {
            "name": "Payment",
            "headers": ["Field", "Type", "Constraint"],
            "data": [
                ["PAYMENT_ID", "VARCHAR", "Not Null"],
                ["BOOKING_ID", "INT", "FK"],
                ["amount", "VARCHAR", "Not Null"],
                ["status", "VARCHAR", "Not Null"]
            ]
        }
    ]
    add_table_slide("Tables: Transactions", tables_s7)

    # Slide 8: Prediction & Chat
    tables_s8 = [
        {
            "name": "SkinPredictionHistory",
            "headers": ["Field", "Type", "Constraint"],
            "data": [
                ["USER_ID", "INT", "FK"],
                ["image_path", "VARCHAR", "Not Null"],
                ["label", "VARCHAR", "Not Null"],
                ["confidence", "FLOAT", "Not Null"]
            ]
        },
        {
            "name": "SkinChatMessage",
            "headers": ["Field", "Type", "Constraint"],
            "data": [
                ["USER_ID", "INT", "FK"],
                ["role", "VARCHAR", "user/bot"],
                ["message", "TEXT", "Not Null"],
                ["created_at", "DATETIME", "Auto"]
            ]
        }
    ]
    add_table_slide("Tables: AI & Chat", tables_s8)

    # 6. DFDs
    add_placeholder_slide("Data Flow Diagram", "Level - 0 (Context Diagram)\n[Insert DFD Level 0 Image Here]")
    add_placeholder_slide("Data Flow Diagram - Admin", "Level - 1 Admin\n[Insert DFD Level 1 Admin Image Here]")
    add_placeholder_slide("Data Flow Diagram - Doctor", "Level - 1 Doctor\n[Insert DFD Level 1 Doctor Image Here]")
    add_placeholder_slide("Data Flow Diagram - Patient", "Level - 1 Patient\n[Insert DFD Level 1 Patient Image Here]")

    # 7. Future Enhancement
    add_bullet_slide("Future Enhancement", [
        "Mobile Application: Developing native iOS/Android apps for on-the-go access.",
        "Telemedicine: Integrating real-time video consultation between doctors and patients.",
        "Blockchain: Implementing blockchain for immutable and secure medical record storage.",
        "Advanced Analytics: Personalized skin care routines based on long-term data analysis."
    ])

    # 8. Screenshots
    add_placeholder_slide("Screenshots: Admin", "Admin Dashboard & User Management\n[Insert Admin Screenshots Here]")
    add_placeholder_slide("Screenshots: Doctor", "Doctor Dashboard & Consultations\n[Insert Doctor Screenshots Here]")
    add_placeholder_slide("Screenshots: Patient", "Patient Prediction & Booking\n[Insert Patient Screenshots Here]")

    # 9. Bibliography
    add_bullet_slide("Bibliography", [
        "Django Documentation: https://docs.djangoproject.com/",
        "Bootstrap 5 Documentation: https://getbootstrap.com/",
        "Python 3 Standard Library: https://docs.python.org/3/library",
        "TensorFlow/Keras API Docs: https://www.tensorflow.org/api_docs",
        "E. Balaguruswamy, 'Introduction to Computing and Problem Solving Using Python'"
    ])

    # Save
    try:
        output_file = 'Dermascan_AI_Presentation_v3.pptx'
        prs.save(output_file)
        print(f"Presentation saved successfully as '{output_file}'")
    except Exception as e:
        print(f"Error saving presentation: {e}")

if __name__ == "__main__":
    create_presentation()
